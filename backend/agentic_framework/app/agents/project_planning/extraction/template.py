"""Generate the student SRS template (.docx) and a sample filled SRS for tests.

Run from backend/agentic_framework/:
    py -m uv run python -m app.agents.project_planning.extraction.template

Outputs:
    frontend/public/SRS_Template.docx                      (student download)
    app/agents/project_planning/extraction/samples/sample_srs.docx  (tests/demo)

The layout matches ``extractor.py``: two tables, headers
``ID | Requirement | Priority`` and ``ID | Requirement | Category``.
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

_FR_HEADERS = ("ID", "Requirement", "Priority")
_NFR_HEADERS = ("ID", "Requirement", "Category")

# Sample SRS content — demonstrating ISO/IEC/IEEE 29148 Score 4 standards alongside contrastive examples.
_SAMPLE_FR = [
    ("FR-01", "The system shall provide a user-friendly login page that is fast and secure.", "High"),
    ("FR-02", "When a registered customer enters invalid credentials 5 consecutive times within 10 minutes, the system shall lock the account for 15 minutes, display the remaining lockout duration, and send a security alert email to the customer's registered email address. Source: Business Rule BR-SEC-01. Test: TC-AUTH-005.", "High"),
    ("FR-03", "When an authenticated student uploads an SRS document in .docx format up to 10 MB, the system shall parse the requirement tables within 15 seconds and display the extracted requirement list. If the uploaded document format is invalid or corrupted, the system shall display a descriptive error message indicating the file formatting fault. Source: Stakeholder Need STK-DOC-02. Test: TC-PARSE-014.", "High"),
    ("FR-04", "The system should handle errors appropriately and notify the relevant people.", "Medium"),
    ("FR-05", "When an instructor reviews a blocked requirement in the quality dashboard, the system shall allow the instructor to apply an administrative override with mandatory justification notes, recording the instructor ID, justification text, and UTC timestamp in the audit log. Source: Governance Rule POL-GOV-03. Test: TC-OVR-021.", "Medium"),
]
_SAMPLE_NFR = [
    ("NFR-01", "Under normal production load up to 100 concurrent requests, the system shall complete and return a quality analysis for a single requirement within 5 seconds in 95% of requests. Source: Performance SLA SLA-PERF-01. Test: TC-PERF-001.", "Performance"),
    ("NFR-02", "The system shall be reliable.", "Reliability"),
]


def _add_table(doc: Document, headers: tuple[str, ...], rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value


def _base_document(title_suffix: str) -> Document:
    doc = Document()
    title = doc.add_heading("Software Requirements Specification (SRS)", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(title_suffix).italic = True

    doc.add_heading("1. Introduction", level=1)
    doc.add_heading("1.1 Purpose", level=2)
    doc.add_paragraph(
        "Describe the purpose of this system in 2–3 sentences. (This section is "
        "not graded by the Quality agent; only the requirement tables below are.)"
    )
    doc.add_heading("1.2 Scope", level=2)
    doc.add_paragraph("Briefly describe what the system does and its boundaries.")

    doc.add_heading("2. Functional Requirements", level=1)
    doc.add_paragraph(
        "List each functional requirement in the table below. Keep one requirement "
        "per row. Follow ISO/IEC/IEEE 29148: specify actor, trigger/precondition, action, "
        "measurable criteria, and exception behavior. Include backward source (Source: ...) "
        "and forward verification (Test: ...) to achieve a Score 4 (Excellent)."
    )
    return doc


def _template_document():
    doc = _base_document("Project: ____________   Team: ____________   Date: ____________")
    empty_fr = [(f"FR-{i:02d}", "", "") for i in range(1, 9)]
    _add_table(doc, _FR_HEADERS, empty_fr)

    doc.add_heading("3. Non-Functional Requirements", level=1)
    doc.add_paragraph("List quality attributes (performance, security, usability, …), one per row.")
    empty_nfr = [(f"NFR-{i:02d}", "", "") for i in range(1, 5)]
    _add_table(doc, _NFR_HEADERS, empty_nfr)

    note = doc.add_paragraph()
    run = note.add_run(
        "ISO/IEC/IEEE 29148 Scoring Tip: To score 4/4 (100% - Excellent), ensure your requirement "
        "specifies: (1) Actor, (2) Trigger or precondition, (3) Measurable action/constraint, "
        "(4) Exception handling behavior, (5) Source link (e.g. Source: BR-xxx / STK-xxx), and "
        "(6) Verification link (e.g. Test: TC-xxx). Avoid vague adjectives like 'fast', 'secure', or 'user-friendly'."
    )
    run.font.size = Pt(9)
    run.italic = True
    return doc


def build_template_bytes() -> bytes:
    """Generate the blank SRS template as .docx bytes (served by the backend)."""
    buffer = io.BytesIO()
    _template_document().save(buffer)
    return buffer.getvalue()


def build_template(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    _template_document().save(str(path))


def build_sample(path: Path) -> None:
    doc = _base_document("Project: Library System   Team: J26-SE-363   Date: 2026-10-13")
    _add_table(doc, _FR_HEADERS, _SAMPLE_FR)
    doc.add_heading("3. Non-Functional Requirements", level=1)
    _add_table(doc, _NFR_HEADERS, _SAMPLE_NFR)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def main() -> None:
    repo_root = Path(__file__).resolve().parents[6]
    template_path = repo_root / "frontend" / "public" / "SRS_Template.docx"
    sample_path = Path(__file__).resolve().parent / "samples" / "sample_srs.docx"
    build_template(template_path)
    build_sample(sample_path)
    print(f"template -> {template_path}")
    print(f"sample   -> {sample_path}")


if __name__ == "__main__":
    main()
