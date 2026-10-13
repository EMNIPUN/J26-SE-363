"""Rule-based requirement extraction from an SRS .docx (the ``extractor`` of the
component diagram).

The companion template (``template.py``) lays requirements out in two tables:

    Functional Requirements      | ID | Requirement | Priority |
    Non-Functional Requirements  | ID | Requirement | Category |

Primary strategy reads those tables. A paragraph fallback (``FR-01: ...`` lines)
makes extraction robust to students who paste requirements as plain text.
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.document import Document as DocxDocument

from app.agents.project_planning.schemas import RequirementInput, RequirementType

# "FR-01", "NFR 2", "fr3:" ... -> capture prefix + number.
_ID_RE = re.compile(r"^\s*(FR|NFR)\s*[-_ ]?\s*(\d+)\s*[:.\)]?\s*(.*)$", re.IGNORECASE)
# Obvious placeholder/guidance text we must not treat as a real requirement.
_PLACEHOLDER_HINTS = (
    "the system shall …",
    "the system shall ...",
    "replace this",
    "e.g.",
    "example:",
    "<",
)


def _req_type(req_id: str) -> RequirementType:
    return (
        RequirementType.NON_FUNCTIONAL
        if req_id.upper().startswith("NFR")
        else RequirementType.FUNCTIONAL
    )


def _normalize_id(prefix: str, number: str) -> str:
    return f"{prefix.upper()}-{int(number):02d}"


def _is_placeholder(text: str) -> bool:
    low = text.strip().lower()
    if not low:
        return True
    return any(low.startswith(h) or low == h for h in _PLACEHOLDER_HINTS)


def _from_tables(doc: DocxDocument) -> list[RequirementInput]:
    found: list[RequirementInput] = []
    for table in doc.tables:
        if not table.rows:
            continue
        header = [c.text.strip().lower() for c in table.rows[0].cells]
        try:
            id_col = next(i for i, h in enumerate(header) if h in ("id", "req id", "ref"))
            req_col = next(i for i, h in enumerate(header) if "requirement" in h or "description" in h)
        except StopIteration:
            continue  # not a requirements table

        for row in table.rows[1:]:
            cells = row.cells
            if max(id_col, req_col) >= len(cells):
                continue
            raw_id = cells[id_col].text.strip()
            text = cells[req_col].text.strip()
            m = _ID_RE.match(raw_id)
            if not m or _is_placeholder(text):
                continue
            req_id = _normalize_id(m.group(1), m.group(2))
            found.append(RequirementInput(id=req_id, text=text, type=_req_type(req_id)))
    return found


def _from_paragraphs(doc: DocxDocument) -> list[RequirementInput]:
    found: list[RequirementInput] = []
    for para in doc.paragraphs:
        m = _ID_RE.match(para.text.strip())
        if not m:
            continue
        text = m.group(3).strip()
        if _is_placeholder(text):
            continue
        req_id = _normalize_id(m.group(1), m.group(2))
        found.append(RequirementInput(id=req_id, text=text, type=_req_type(req_id)))
    return found


def _dedupe(reqs: list[RequirementInput]) -> list[RequirementInput]:
    seen: dict[str, RequirementInput] = {}
    for r in reqs:
        seen.setdefault(r.id, r)  # first occurrence wins
    return list(seen.values())


def extract_from_document(doc: DocxDocument) -> list[RequirementInput]:
    reqs = _from_tables(doc)
    if not reqs:  # fall back only if the tables yielded nothing
        reqs = _from_paragraphs(doc)
    return _dedupe(reqs)


def extract_from_bytes(data: bytes) -> list[RequirementInput]:
    return extract_from_document(Document(io.BytesIO(data)))


def extract_from_path(path: str) -> list[RequirementInput]:
    return extract_from_document(Document(path))
